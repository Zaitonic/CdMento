import io
import hashlib
import pytest
from fastapi import Request
from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter
from backend.config import Settings
from backend.repository import Repository
from backend.storage import FileStorage
from backend.services import Services
from backend.ai import AIService
from backend.main import create_app, user
from backend.documents import extract


class FakeAI:
    """Deterministic test double; tests never transmit source material to a provider."""

    client = None

    def require(self):
        pass

    def answer(self, notebook, sources, question, history):
        return {
            "answer": "A LAN connects nearby devices.",
            "source_ids": [sources[0]["id"]],
        }

    def generate(self, request, sources):
        sid = sources[0]["id"]
        if request["kind"] == "flashcards":
            return {
                "cards": [
                    {
                        "question": "What is a LAN?",
                        "answer": "A local network.",
                        "concept": "LAN",
                        "source_ids": [sid],
                    }
                ]
            }
        if request["kind"] == "quiz":
            return {
                "questions": [
                    {
                        "question": "Which network connects nearby devices?",
                        "options": ["LAN", "WAN", "DNS", "TCP"],
                        "correct": 0,
                        "explanation": "A LAN connects devices within a local area.",
                        "concept": "LAN",
                        "source_ids": [sid],
                    }
                ]
            }
        return {
            "nodes": [{"id": "lan", "label": "LAN", "source_ids": [sid]}],
            "edges": [],
        }


@pytest.fixture
def env(tmp_path):
    config = Settings(
        _env_file=None,
        data_dir=tmp_path,
        worker_enabled=False,
        openai_api_key="",
        auth_mode="local",
        data_backend="sqlite",
    )
    services = Services(config, Repository(tmp_path), FileStorage(tmp_path), FakeAI())
    services.repo.put(
        "alice", "preferences", {"id": "ai", "providers": ["groq"]}
    )
    app = create_app(services=services)

    def test_identity(request: Request):
        return request.headers.get("X-Test-User", "alice")

    app.dependency_overrides[user] = test_identity
    with TestClient(app) as client:
        yield client, services, app


def seed(client, services):
    n = client.post("/api/v1/notebooks", json={"name": "Networking"}).json()
    result = client.post(
        f"/api/v1/notebooks/{n['id']}/sources",
        files={
            "file": (
                "chapter.txt",
                b"A LAN connects nearby devices. A WAN spans wide geographic areas.",
                "text/plain",
            )
        },
    ).json()
    services.run_job("alice", result["job_id"])
    return n, result["source"]["id"]


def generated(client, services, nid, kind):
    result = client.post(
        f"/api/v1/notebooks/{nid}/generations", json={"kind": kind}
    ).json()
    services.run_job("alice", result["job_id"])
    return client.get("/api/v1/artifacts/" + result["artifact_id"]).json()


def test_new_workspace_empty_and_notebook_persists(env):
    client, services, _ = env
    assert client.get("/api/v1/workspace").json() == {
        "notebooks": [],
        "sources": [],
        "notes": [],
        "artifacts": [],
        "conversations": [],
        "focus_sessions": [],
        "tasks": [],
        "reviews": [],
        "attempts": [],
        "maps": [],
    }
    n = client.post("/api/v1/notebooks", json={"name": "  Biology  "}).json()
    assert n["name"] == "Biology"
    reopened = Repository(services.settings.data_dir)
    assert reopened.get("alice", "notebooks", n["id"])["name"] == "Biology"
    assert client.post("/api/v1/notebooks", json={"name": "   "}).status_code == 422


def test_owner_isolation_across_data_and_jobs(env):
    client, services, _ = env
    n, sid = seed(client, services)
    other = {"X-Test-User": "bob"}
    assert client.get("/api/v1/workspace", headers=other).json()["notebooks"] == []
    assert (
        client.get(f"/api/v1/sources/{sid}/preview", headers=other).status_code == 404
    )
    assert (
        client.patch(
            f"/api/v1/notebooks/{n['id']}", json={"name": "Stolen"}, headers=other
        ).status_code
        == 404
    )
    assert (
        client.post(
            f"/api/v1/notebooks/{n['id']}/notes",
            json={"title": "Stolen"},
            headers=other,
        ).status_code
        == 404
    )
    assert (
        client.get(
            "/api/v1/jobs/" + services.repo.list("alice", "jobs")[0]["id"],
            headers=other,
        ).status_code
        == 404
    )


def test_note_conflict_preserves_newer_document(env):
    client, _, _ = env
    n = client.post("/api/v1/notebooks", json={"name": "Notes"}).json()
    note = client.post(
        f"/api/v1/notebooks/{n['id']}/notes", json={"title": "Draft"}
    ).json()
    update = {
        "title": "Updated",
        "content": "<p>New</p>",
        "document": {"type": "doc", "content": []},
        "revision": note["revision"],
    }
    assert client.patch("/api/v1/notes/" + note["id"], json=update).status_code == 200
    assert (
        client.patch(
            "/api/v1/notes/" + note["id"], json={**update, "title": "Old edit"}
        ).status_code
        == 409
    )
    assert client.get("/api/v1/workspace").json()["notes"][0]["title"] == "Updated"


def test_upload_extract_preview_original_and_duplicates(env):
    client, services, _ = env
    n, sid = seed(client, services)
    preview = client.get(f"/api/v1/sources/{sid}/preview").json()
    assert preview["status"] == "ready" and "LAN" in preview["text"]
    source = services.repo.get("alice", "sources", sid)
    assert services.storage.get(source["storage_key"]).startswith(b"A LAN")
    repeat = client.post(
        f"/api/v1/notebooks/{n['id']}/sources",
        files={
            "file": (
                "again.txt",
                services.storage.get(source["storage_key"]),
                "text/plain",
            )
        },
    ).json()
    assert repeat["duplicate"] is True
    assert len(client.get("/api/v1/workspace").json()["sources"]) == 1


def test_invalid_upload_and_pdf_needing_ocr(env):
    client, services, _ = env
    n = client.post("/api/v1/notebooks", json={"name": "Files"}).json()
    path = f"/api/v1/notebooks/{n['id']}/sources"
    assert client.post(path, files={"file": ("bad.exe", b"text")}).status_code == 415
    assert client.post(path, files={"file": ("empty.txt", b"")}).status_code == 400
    bad = client.post(path, files={"file": ("fake.pdf", b"not a pdf")}).json()
    services.run_job("alice", bad["job_id"])
    assert client.get("/api/v1/jobs/" + bad["job_id"]).json()["status"] == "failed"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    blank = client.post(path, files={"file": ("scan.pdf", buffer.getvalue())}).json()
    services.run_job("alice", blank["job_id"])
    assert (
        client.get("/api/v1/sources/" + blank["source"]["id"] + "/preview").json()[
            "status"
        ]
        == "needs_ocr"
    )
    oversized = client.post(
        path, files={"file": ("huge.txt", b"x" * (10 * 1024 * 1024 + 1))}
    )
    assert oversized.status_code == 413


def test_docx_tables_and_encoding():
    document = Document()
    document.add_paragraph("A network connects devices.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "LAN"
    table.cell(0, 1).text = "Local area network"
    buffer = io.BytesIO()
    document.save(buffer)
    text, chunks = extract(buffer.getvalue(), "docx")
    assert "Local area network" in text and len(chunks) == 2
    assert extract("Hello world".encode("utf-16"), "txt")[0] == "Hello world"


def test_chat_uses_owned_sources_and_persists(env):
    client, services, _ = env
    n, sid = seed(client, services)
    path = f"/api/v1/notebooks/{n['id']}/messages"
    assert (
        client.post(path, json={"question": "LAN?", "source_ids": []}).status_code
        == 400
    )
    assert (
        client.post(
            path, json={"question": "LAN?", "source_ids": ["foreign-id"]}
        ).status_code
        == 404
    )
    answer = client.post(path, json={"question": "LAN?", "source_ids": [sid]}).json()
    assert answer["citations"][0]["source_id"] == sid
    assert (
        len(client.get(f"/api/v1/notebooks/{n['id']}/conversation").json()["messages"])
        == 2
    )
    services.ai = AIService(services.settings)
    assert client.post(path, json={"question": "LAN?"}).status_code == 503
    assert (
        client.post(
            path, json={"question": "LAN?", "context": "Injected document"}
        ).status_code
        == 422
    )


def test_generation_cache_card_acceptance_and_fsrs_idempotence(env):
    client, services, _ = env
    n, _ = seed(client, services)
    deck = generated(client, services, n["id"], "flashcards")
    cached = client.post(
        f"/api/v1/notebooks/{n['id']}/generations", json={"kind": "flashcards"}
    ).json()
    assert cached["cached"] is True and cached["artifact_id"] == deck["id"]
    card = deck["cards"][0]
    path = f"/api/v1/decks/{deck['id']}/cards/{card['id']}"
    assert (
        client.post(
            path + "/reviews", json={"rating": 3, "event_id": "review-1"}
        ).status_code
        == 409
    )
    assert (
        client.patch(
            path,
            json={
                "question": card["question"],
                "answer": card["answer"],
                "accepted": True,
            },
        ).status_code
        == 200
    )
    result = client.post(
        path + "/reviews", json={"rating": 3, "event_id": "review-1"}
    ).json()
    repeat = client.post(
        path + "/reviews", json={"rating": 3, "event_id": "review-1"}
    ).json()
    assert result["due"] == repeat["due"]
    assert len(services.repo.list("alice", "reviews")) == 1


def test_quiz_answer_keys_hidden_server_grades_and_retries_safe(env):
    client, services, _ = env
    n, _ = seed(client, services)
    quiz = generated(client, services, n["id"], "quiz")
    assert (
        "correct" not in quiz["questions"][0]
        and "explanation" not in quiz["questions"][0]
    )
    attempt = client.post(
        f"/api/v1/quizzes/{quiz['id']}/attempts", json={"mode": "exam"}
    ).json()
    q = attempt["questions"][0]
    assert "correct" not in q
    assert (
        client.post(
            f"/api/v1/attempts/{attempt['id']}/check",
            json={"question_id": q["id"], "option_id": "0"},
        ).status_code
        == 403
    )
    path = f"/api/v1/attempts/{attempt['id']}/submit"
    assert (
        client.post(path, json={"answers": {q["id"]: "not-an-option"}}).status_code
        == 422
    )
    result = client.post(path, json={"answers": {q["id"]: "1"}}).json()
    assert result["score"] == 0 and result["weak_concepts"] == ["LAN"]
    assert client.post(path, json={"answers": {q["id"]: "0"}}).json() == result


def test_maps_focus_and_tasks_persist(env):
    client, services, _ = env
    n, _ = seed(client, services)
    path = f"/api/v1/notebooks/{n['id']}/mind-map"
    assert (
        client.put(
            path,
            json={
                "nodes": [{"id": "a"}],
                "edges": [{"source": "a", "target": "missing"}],
            },
        ).status_code
        == 422
    )
    graph = {
        "nodes": [{"id": "a", "position": {"x": 1, "y": 2}, "data": {"label": "LAN"}}],
        "edges": [],
        "revision": 0,
    }
    assert client.put(path, json=graph).status_code == 200
    assert client.put(path, json=graph).status_code == 409
    focus = client.post("/api/v1/focus-sessions", json={"duration": 1500}).json()
    assert (
        client.post("/api/v1/focus-sessions", json={"duration": 300}).status_code == 409
    )
    assert (
        client.patch(
            "/api/v1/focus-sessions/" + focus["id"], json={"action": "complete"}
        ).status_code
        == 409
    )
    paused = client.patch(
        "/api/v1/focus-sessions/" + focus["id"], json={"action": "pause"}
    ).json()
    assert paused["status"] == "paused" and paused["remaining"] > 0
    assert client.get("/api/v1/focus-sessions/current").json()["id"] == focus["id"]
    task = client.post(
        "/api/v1/tasks", json={"text": "Review chapter", "done": False}
    ).json()
    assert (
        client.patch(
            "/api/v1/tasks/" + task["id"], json={"text": task["text"], "done": True}
        ).json()["done"]
        is True
    )


def test_deleting_notebook_cleans_originals_and_related_records(env):
    client, services, _ = env
    n, sid = seed(client, services)
    source = services.repo.get("alice", "sources", sid)
    client.post(f"/api/v1/notebooks/{n['id']}/notes", json={"title": "Saved"})
    job = client.delete("/api/v1/notebooks/" + n["id"]).json()
    assert client.get("/api/v1/workspace").json()["notebooks"] == []
    services.run_job("alice", job["id"])
    assert (
        services.repo.list("alice", "sources") == []
        and services.repo.list("alice", "notes") == []
    )
    assert not services.storage.local_path(source["storage_key"]).exists()


def test_import_idempotent_and_unrelated_work_preserved(env):
    client, _, _ = env
    original = client.post("/api/v1/notebooks", json={"name": "Existing"}).json()
    backup = {
        "notebooks": [{"id": "legacy", "name": "Imported", "color": "purple"}],
        "notes": [
            {
                "id": "note1",
                "notebook": "legacy",
                "title": "Old note",
                "content": "<p>Preserved</p>",
            }
        ],
        "sources": [],
    }
    assert client.post("/api/v1/workspace/import", json=backup).status_code == 200
    result = client.post("/api/v1/workspace/import", json=backup).json()
    assert len(result["notebooks"]) == 2 and len(result["notes"]) == 1
    assert any(n["id"] == original["id"] for n in result["notebooks"])


def test_firebase_missing_invalid_and_valid_tokens(env, monkeypatch):
    client, services, app = env
    services.settings.auth_mode = "firebase"
    app.dependency_overrides.clear()
    from firebase_admin import auth

    def verify(token, **kwargs):
        if token != "valid":
            raise ValueError("Bad token")
        return {"uid": "firebase-user", "email_verified": True}

    monkeypatch.setattr(auth, "verify_id_token", verify)
    assert client.get("/api/v1/me").status_code == 401
    assert (
        client.get(
            "/api/v1/me", headers={"Authorization": "Bearer expired"}
        ).status_code
        == 401
    )
    assert (
        client.get("/api/v1/me", headers={"Authorization": "Bearer valid"}).json()["id"]
        == "firebase-user"
    )


def test_local_mode_rejects_cross_origin(env):
    client, _, _ = env
    assert (
        client.get(
            "/api/v1/health", headers={"Origin": "https://untrusted.example"}
        ).status_code
        == 403
    )


def test_local_demo_has_isolated_workspace_and_revocable_session(env):
    client, services, app = env
    app.dependency_overrides.clear()
    assert client.get("/api/v1/health").json()["demo_account_available"] is False
    local = client.post(
        "/api/v1/notebooks", json={"name": "Original local work"}
    ).json()
    app.state.demo_auth.create_account("demo@example.com", "test-only-password")
    assert client.get("/api/v1/health").json()["demo_account_available"] is True
    assert (
        client.post(
            "/api/v1/auth/demo/login",
            json={"email": "demo@example.com", "password": "incorrect"},
        ).status_code
        == 401
    )
    result = client.post(
        "/api/v1/auth/demo/login",
        json={"email": "DEMO@example.com", "password": "test-only-password"},
    )
    assert result.status_code == 200
    data = result.json()
    assert data["profile"]["kind"] == "local-demo"
    assert "password_hash" not in data["profile"]
    headers = {"Authorization": "Bearer " + data["token"]}
    assert (
        client.get("/api/v1/auth/demo/session", headers=headers).json()["uid"]
        == "demo-student"
    )
    assert client.get("/api/v1/workspace", headers=headers).json()["notebooks"] == []
    demo = client.post(
        "/api/v1/notebooks", json={"name": "Demo work"}, headers=headers
    ).json()
    assert (
        client.get("/api/v1/workspace", headers=headers).json()["notebooks"][0]["id"]
        == demo["id"]
    )
    assert client.get("/api/v1/workspace").json()["notebooks"][0]["id"] == local["id"]
    from backend.demo_auth import DemoAuth

    assert (
        DemoAuth(Repository(services.settings.data_dir)).resolve(data["token"])
        == "demo-student"
    )
    assert client.post("/api/v1/auth/demo/logout", headers=headers).status_code == 204
    assert client.get("/api/v1/workspace", headers=headers).status_code == 401


def test_expired_demo_session_does_not_fall_back_to_local_workspace(env):
    client, services, app = env
    app.dependency_overrides.clear()
    app.state.demo_auth.create_account("demo@example.com", "test-only-password")
    token = client.post(
        "/api/v1/auth/demo/login",
        json={"email": "demo@example.com", "password": "test-only-password"},
    ).json()["token"]
    from backend.demo_auth import AUTH_OWNER

    key = hashlib.sha256(token.encode()).hexdigest()
    session = services.repo.get(AUTH_OWNER, "sessions", key)
    services.repo.put(AUTH_OWNER, "sessions", {**session, "expires_at": 0})
    headers = {"Authorization": "Bearer " + token}
    assert client.get("/api/v1/workspace", headers=headers).status_code == 401
    assert (
        client.post(
            "/api/v1/notebooks", json={"name": "Must not save"}, headers=headers
        ).status_code
        == 401
    )
    assert services.repo.list("local-workspace", "notebooks") == []


def test_demo_signin_is_rate_limited_and_disabled_in_firebase_mode(env):
    client, services, app = env
    app.state.demo_auth.create_account("demo@example.com", "test-only-password")
    credentials = {"email": "unknown@example.com", "password": "incorrect"}
    for _ in range(10):
        assert (
            client.post("/api/v1/auth/demo/login", json=credentials).status_code == 401
        )
    assert client.post("/api/v1/auth/demo/login", json=credentials).status_code == 429
    services.settings.auth_mode = "firebase"
    assert client.post("/api/v1/auth/demo/login", json=credentials).status_code == 404
    assert client.get("/api/v1/auth/demo/session").status_code == 404
    assert client.get("/api/v1/health").json()["demo_account_available"] is False


def test_unverified_firebase_accounts_cannot_open_workspace(env, monkeypatch):
    client, services, app = env
    from firebase_admin import auth

    services.settings.auth_mode = "firebase"
    app.dependency_overrides.clear()
    monkeypatch.setattr(
        auth,
        "verify_id_token",
        lambda *args, **kwargs: {
            "uid": "unverified",
            "email_verified": False,
        },
    )
    headers = {"Authorization": "Bearer unverified"}
    assert client.get("/api/v1/workspace", headers=headers).status_code == 403
    assert (
        client.post(
            "/api/v1/notebooks", json={"name": "Blocked"}, headers=headers
        ).status_code
        == 403
    )
    assert client.put("/api/v1/me/profile", headers=headers).status_code == 403
    assert services.repo.list("unverified", "notebooks") == []


def test_account_profile_uses_verified_claims_and_is_owner_scoped(env, monkeypatch):
    client, services, app = env
    from firebase_admin import auth

    services.settings.auth_mode = "firebase"
    app.dependency_overrides.clear()
    monkeypatch.setattr(
        auth,
        "verify_id_token",
        lambda token, **kwargs: {
            "uid": token,
            "email_verified": True,
            "email": token + "@example.com",
            "name": "Verified " + token,
        },
    )
    headers = {"Authorization": "Bearer alice"}
    first = client.put(
        "/api/v1/me/profile",
        headers=headers,
        json={
            "email": "forged@example.com",
            "display_name": "Forged",
            "role": "admin",
        },
    )
    assert first.status_code == 200
    assert first.json()["email"] == "alice@example.com"
    assert first.json()["display_name"] == "Verified alice"
    assert "role" not in first.json()
    second = client.put("/api/v1/me/profile", headers=headers).json()
    assert second["created_at"] == first.json()["created_at"]
    edited = client.patch(
        "/api/v1/me/profile",
        headers=headers,
        json={"display_name": "My chosen name", "bio": "Learning every day"},
    )
    assert edited.status_code == 200
    reopened = client.put("/api/v1/me/profile", headers=headers).json()
    assert reopened["display_name"] == "My chosen name"
    assert reopened["bio"] == "Learning every day"
    assert reopened["email"] == "alice@example.com"
    client.put("/api/v1/me/profile", headers={"Authorization": "Bearer bob"})
    assert (
        services.repo.get("alice", "profile", "account")["email"] == "alice@example.com"
    )
    assert services.repo.get("bob", "profile", "account")["email"] == "bob@example.com"


def test_queued_upload_survives_repository_reopen(env):
    client, services, _ = env
    n = client.post("/api/v1/notebooks", json={"name": "Restart check"}).json()
    job = client.post(
        f"/api/v1/notebooks/{n['id']}/sources",
        files={"file": ("chapter.txt", b"Queued source survives a restart.")},
    ).json()
    restarted = Services(
        services.settings,
        Repository(services.settings.data_dir),
        services.storage,
        FakeAI(),
    )
    restarted.run_job("alice", job["job_id"])
    assert client.get("/api/v1/jobs/" + job["job_id"]).json()["status"] == "completed"
    assert (
        "survives"
        in client.get("/api/v1/sources/" + job["source"]["id"] + "/preview").json()[
            "text"
        ]
    )


def test_invalid_import_is_rejected_before_creating_notebooks(env):
    client, _, _ = env
    backup = {
        "notebooks": [{"id": "legacy", "name": "Would be imported"}],
        "sources": [
            {"id": "source", "notebook": "legacy", "name": "empty.txt", "text": ""}
        ],
        "notes": [],
    }
    assert client.post("/api/v1/workspace/import", json=backup).status_code == 422
    assert client.get("/api/v1/workspace").json()["notebooks"] == []


def test_malformed_json_and_invalid_generation_are_json_errors(env):
    client, services, _ = env
    assert (
        client.post(
            "/api/v1/notebooks",
            content="{invalid",
            headers={"Content-Type": "application/json"},
        ).status_code
        == 422
    )
    n, _ = seed(client, services)
    assert (
        client.post(
            f"/api/v1/notebooks/{n['id']}/generations",
            json={"kind": "quiz", "count": 999},
        ).status_code
        == 422
    )


def test_source_original_is_protected_and_retains_bytes(env):
    client, services, _ = env
    notebook, sid = seed(client, services)
    result = client.get(f"/api/v1/sources/{sid}/file")
    assert result.status_code == 200
    assert (
        result.content
        == b"A LAN connects nearby devices. A WAN spans wide geographic areas."
    )
    assert result.headers["cache-control"] == "no-store"
    assert result.headers["x-content-type-options"] == "nosniff"
    assert (
        client.get(
            f"/api/v1/sources/{sid}/file", headers={"X-Test-User": "bob"}
        ).status_code
        == 404
    )
    assert client.get("/api/v1/sources/missing/file").status_code == 404


def test_pdf_original_can_be_viewed_when_ocr_is_needed(env):
    client, services, _ = env
    notebook = client.post("/api/v1/notebooks", json={"name": "Scanned notes"}).json()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    output = io.BytesIO()
    writer.write(output)
    content = output.getvalue()
    upload = client.post(
        f"/api/v1/notebooks/{notebook['id']}/sources",
        files={"file": ("scan.pdf", content, "application/pdf")},
    ).json()
    services.run_job("alice", upload["job_id"])
    sid = upload["source"]["id"]
    assert client.get(f"/api/v1/sources/{sid}/preview").json()["status"] == "needs_ocr"
    original = client.get(f"/api/v1/sources/{sid}/file")
    assert original.content == content
    assert original.headers["content-type"] == "application/pdf"


def test_partially_readable_pdf_reports_missing_pages(env):
    from pypdf.generic import DecodedStreamObject, NameObject, DictionaryObject

    client, services, _ = env
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 20 200 Td (JavaScript variables store values.) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_blank_page(width=300, height=300)
    output = io.BytesIO()
    writer.write(output)
    notebook = client.post("/api/v1/notebooks", json={"name": "Programming"}).json()
    upload = client.post(
        f"/api/v1/notebooks/{notebook['id']}/sources",
        files={"file": ("lesson.pdf", output.getvalue(), "application/pdf")},
    ).json()
    services.run_job("alice", upload["job_id"])
    preview = client.get(f"/api/v1/sources/{upload['source']['id']}/preview").json()
    assert preview["status"] == "ready"
    assert preview["page_count"] == 2
    assert preview["unreadable_pages"] == [2]
    assert "variables store values" in preview["text"]


def test_reading_sections_preserves_unicode_across_overlapping_chunks():
    from backend.documents import reading_sections

    content = "Study 🧠 " * 800
    _, chunks = extract(content.encode("utf-8"), "txt")
    sections = reading_sections(chunks)
    assert len(sections) == 1
    assert sections[0]["text"] == content.strip()


def test_reprocess_existing_source_is_owner_scoped_and_deduplicated(env):
    client, services, _ = env
    _, sid = seed(client, services)
    assert (
        client.post(
            f"/api/v1/sources/{sid}/reprocess", headers={"X-Test-User": "bob"}
        ).status_code
        == 404
    )
    first = client.post(f"/api/v1/sources/{sid}/reprocess")
    second = client.post(f"/api/v1/sources/{sid}/reprocess")
    assert first.status_code == 202 and first.json() == second.json()
    services.run_job("alice", first.json()["job_id"])
    assert client.get(f"/api/v1/sources/{sid}/preview").json()["status"] == "ready"


def test_large_original_stream_retains_final_bytes(env):
    client, _, _ = env
    notebook = client.post("/api/v1/notebooks", json={"name": "Large original"}).json()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    output = io.BytesIO()
    writer.write(output)
    content = output.getvalue() + b"\n%" + b"x" * 3_397_423 + b"\n"
    upload = client.post(
        f"/api/v1/notebooks/{notebook['id']}/sources",
        files={"file": ("large.pdf", content, "application/pdf")},
    ).json()
    response = client.get(f"/api/v1/sources/{upload['source']['id']}/file")
    assert response.status_code == 200 and response.content == content
    assert "content-length" not in response.headers


def test_custom_profile_is_owned_validated_and_persistent(env):
    client, services, _ = env
    profile = {
        "display_name": "  Curious Student  ",
        "bio": "Learning Python",
        "study_goal": "Practice daily",
        "avatar": "plant",
        "color": "green",
    }
    response = client.patch("/api/v1/me/profile", json=profile)
    assert response.status_code == 200
    assert response.json()["display_name"] == "Curious Student"
    assert client.get("/api/v1/me/profile").json()["study_goal"] == "Practice daily"
    bob = client.get("/api/v1/me/profile", headers={"X-Test-User": "bob"}).json()
    assert bob["bio"] == ""
    assert bob["display_name"] != "Curious Student"
    assert (
        Repository(services.settings.data_dir).get("alice", "profile", "account")[
            "avatar"
        ]
        == "plant"
    )
    for changes in [
        {"display_name": "   "},
        {"bio": "x" * 301},
        {"avatar": "https://tracking.example/a"},
        {"email": "other@example.com"},
        {"role": "admin"},
    ]:
        assert (
            client.patch("/api/v1/me/profile", json={**profile, **changes}).status_code
            == 422
        )
    assert client.get("/api/v1/me/profile").json()["display_name"] == "Curious Student"
